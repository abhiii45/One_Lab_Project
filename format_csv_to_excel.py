import subprocess
import sys
from pathlib import Path

# Ensure openpyxl is installed
try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
except ImportError:
    print("Installing openpyxl...")
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl', '-q'])
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

import pandas as pd

# Add reconciliation module to path
PROJECT_ROOT = Path(__file__).resolve().parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


def format_mismatch_report_excel(csv_path, output_path):
    """Create a beautifully formatted Excel file from the mismatch report CSV."""
    
    # Read the CSV
    df = pd.read_csv(csv_path)
    
    # Create Excel writer
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Mismatches', index=False)
        
        # Get the workbook and worksheet
        workbook = writer.book
        worksheet = writer.sheets['Mismatches']
        
        # Define colors for each issue type
        colors = {
            'amount_mismatch': 'FFB6C1',  # Light red
            'cross_month_settlement': 'FFE4B5',  # Moccasin
            'duplicate_settlement': 'FFD700',  # Gold
            'duplicate_transaction': 'FFA500',  # Orange
            'extra_settlement': 'ADD8E6',  # Light blue
            'missing_settlement': 'FFB6C1',  # Light red
            'refund_without_original': 'E6E6FA',  # Lavender
            'rounding_difference': 'F0E68C',  # Khaki
        }
        
        # Set column widths
        column_widths = {
            'A': 25,  # issue_type
            'B': 15,  # transaction_id
            'C': 18,  # transaction_amount
            'D': 18,  # settled_amount
            'E': 16,  # transaction_date
            'F': 16,  # settlement_date
            'G': 60,  # explanation
        }
        
        for col, width in column_widths.items():
            worksheet.column_dimensions[col].width = width
        
        # Style header row
        header_fill = PatternFill(start_color='2F4F7F', end_color='2F4F7F', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        
        # Style data rows
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, max_row=len(df) + 1), start=2):
            # Get issue type for this row
            issue_type = df.iloc[row_idx - 2]['issue_type']
            row_color = colors.get(issue_type, 'FFFFFF')
            row_fill = PatternFill(start_color=row_color, end_color=row_color, fill_type='solid')
            
            for cell in row:
                cell.fill = row_fill
                cell.border = thin_border
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                
                # Format numeric columns
                if cell.column in [3, 4]:  # transaction_amount, settled_amount
                    try:
                        cell.value = float(cell.value) if cell.value and pd.notna(cell.value) else cell.value
                        cell.number_format = '$#,##0.00'
                    except:
                        pass
        
        # Add borders to header
        for cell in worksheet[1]:
            cell.border = thin_border
        
        # Freeze header row
        worksheet.freeze_panes = 'A2'
        
        # Set row height for header
        worksheet.row_dimensions[1].height = 30
        
        # Auto-adjust row heights for wrapped text
        for row_idx in range(2, len(df) + 2):
            worksheet.row_dimensions[row_idx].height = None  # Auto
    
    return output_path


if __name__ == '__main__':
    csv_file = Path('outputs/detailed_mismatch_report.csv')
    excel_file = Path('outputs/detailed_mismatch_report.xlsx')
    
    if csv_file.exists():
        print(f"Converting {csv_file} to Excel format...")
        format_mismatch_report_excel(csv_file, excel_file)
        print(f"✅ Created beautifully formatted Excel file: {excel_file}")
        print(f"📊 File size: {excel_file.stat().st_size / 1024:.1f} KB")
    else:
        print(f"❌ CSV file not found: {csv_file}")
